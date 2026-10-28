"""
Excel Executive Report Generator
Creates a professionally styled, multi-sheet Excel workbook with formatted KPI summaries,
category breakdowns, regional performance tables, and monthly trend analysis.
"""

import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def create_executive_excel():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    clean_csv = os.path.join(base_dir, "data", "processed", "clean_sales_data_complete.csv")
    output_excel = os.path.join(base_dir, "data", "processed", "sales_executive_summary.xlsx")

    df = pd.read_csv(clean_csv)

    # 1. Category aggregation
    cat_summary = df.groupby(["Category", "Sub_Category"]).agg(
        Total_Orders=("Order_ID", "nunique"),
        Total_Units=("Quantity", "sum"),
        Total_Revenue=("Sales", "sum"),
        Total_Profit=("Profit", "sum")
    ).reset_index()
    cat_summary["Profit_Margin_Pct"] = (cat_summary["Total_Profit"] / cat_summary["Total_Revenue"]) * 100
    cat_summary = cat_summary.sort_values(by="Total_Revenue", ascending=False)

    # 2. Regional aggregation
    reg_summary = df.groupby(["Region", "State"]).agg(
        Total_Orders=("Order_ID", "nunique"),
        Total_Units=("Quantity", "sum"),
        Total_Revenue=("Sales", "sum"),
        Total_Profit=("Profit", "sum"),
        Avg_Delivery_Days=("Delivery_Days", "mean")
    ).reset_index()
    reg_summary["Profit_Margin_Pct"] = (reg_summary["Total_Profit"] / reg_summary["Total_Revenue"]) * 100
    reg_summary = reg_summary.sort_values(by="Total_Revenue", ascending=False)

    # 3. Monthly Trend aggregation
    monthly_summary = df.groupby(["Order_Year_Month"]).agg(
        Orders=("Order_ID", "nunique"),
        Units=("Quantity", "sum"),
        Revenue=("Sales", "sum"),
        Profit=("Profit", "sum")
    ).reset_index()
    monthly_summary["Profit_Margin_Pct"] = (monthly_summary["Profit"] / monthly_summary["Revenue"]) * 100
    monthly_summary = monthly_summary.sort_values(by="Order_Year_Month")

    # 4. Overall KPI Card values
    total_rev = df["Sales"].sum()
    total_prof = df["Profit"].sum()
    overall_margin = (total_prof / total_rev) * 100
    total_orders = df["Order_ID"].nunique()
    avg_order_value = total_rev / total_orders
    total_units = df["Quantity"].sum()

    kpi_df = pd.DataFrame([
        {"Key Performance Indicator": "Total Revenue", "Value": total_rev, "Format": "Currency"},
        {"Key Performance Indicator": "Total Profit", "Value": total_prof, "Format": "Currency"},
        {"Key Performance Indicator": "Overall Profit Margin", "Value": overall_margin / 100, "Format": "Percent"},
        {"Key Performance Indicator": "Total Orders Placed", "Value": total_orders, "Format": "Integer"},
        {"Key Performance Indicator": "Average Order Value (AOV)", "Value": avg_order_value, "Format": "Currency"},
        {"Key Performance Indicator": "Total Units Sold", "Value": total_units, "Format": "Integer"},
        {"Key Performance Indicator": "Average Delivery Time (Days)", "Value": df["Delivery_Days"].mean(), "Format": "Decimal"},
    ])

    # Write to Excel with openpyxl
    with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
        kpi_df.to_excel(writer, sheet_name="KPI_Overview", index=False)
        cat_summary.to_excel(writer, sheet_name="Category_Analysis", index=False)
        reg_summary.to_excel(writer, sheet_name="Regional_Analysis", index=False)
        monthly_summary.to_excel(writer, sheet_name="Monthly_Trends", index=False)
        df.head(2000).to_excel(writer, sheet_name="Sample_Clean_Data", index=False)

    # Style the workbook with Corporate Executive Theme
    from openpyxl import load_workbook
    wb = load_workbook(output_excel)

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    kpi_title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for ws in wb.worksheets:
        ws.views.sheetView[0].showGridLines = True

        # Header styling
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 26

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Custom formatting for KPI_Overview
    ws_kpi = wb["KPI_Overview"]
    for row in range(2, len(kpi_df) + 2):
        fmt_type = ws_kpi.cell(row=row, column=3).value
        val_cell = ws_kpi.cell(row=row, column=2)
        if fmt_type == "Currency":
            val_cell.number_format = '$#,##0.00'
        elif fmt_type == "Percent":
            val_cell.number_format = '0.00%'
        elif fmt_type == "Integer":
            val_cell.number_format = '#,##0'
        elif fmt_type == "Decimal":
            val_cell.number_format = '0.0'
        val_cell.alignment = Alignment(horizontal="right")
        val_cell.font = Font(bold=True)

    # Custom formatting for Category Analysis
    ws_cat = wb["Category_Analysis"]
    for row in range(2, len(cat_summary) + 2):
        ws_cat.cell(row=row, column=3).number_format = '#,##0'
        ws_cat.cell(row=row, column=4).number_format = '#,##0'
        ws_cat.cell(row=row, column=5).number_format = '$#,##0.00'
        ws_cat.cell(row=row, column=6).number_format = '$#,##0.00'
        ws_cat.cell(row=row, column=7).number_format = '0.00"%"'

    # Custom formatting for Regional Analysis
    ws_reg = wb["Regional_Analysis"]
    for row in range(2, len(reg_summary) + 2):
        ws_reg.cell(row=row, column=3).number_format = '#,##0'
        ws_reg.cell(row=row, column=4).number_format = '#,##0'
        ws_reg.cell(row=row, column=5).number_format = '$#,##0.00'
        ws_reg.cell(row=row, column=6).number_format = '$#,##0.00'
        ws_reg.cell(row=row, column=7).number_format = '0.1'
        ws_reg.cell(row=row, column=8).number_format = '0.00"%"'

    wb.save(output_excel)
    print(f"Generated formatted Executive Excel Report at: {output_excel}")


if __name__ == "__main__":
    create_executive_excel()
